import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference

def generate_report(input_path, output_path):
    # Step 1: Read the CSV and clean headers
    df = pd.read_csv(input_path)
    df.columns = df.columns.str.strip()

    # Step 2: Clean missing values
    df['Name'] = df['Name'].fillna("Unknown")
    df['Age'] = df['Age'].fillna(df['Age'].mean())
    df['Age'] = df['Age'].astype(int)

    # Step 3: Calculate Tax (10% of Salary)
    df['Tax'] = df['Salary'] * 0.10

    # Step 4: Save to Excel
    df.to_excel(output_path, index=False)

    # Step 5: Add Bar Chart (Salary vs Tax)
    wb = load_workbook(output_path)
    ws = wb.active

    chart = BarChart()
    chart.title = "Salary vs Tax"
    chart.x_axis.title = "Name"
    chart.y_axis.title = "Amount"

    # Add Salary and Tax as series
    data = Reference(ws, min_col=4, max_col=5, min_row=1, max_row=len(df)+1)
    cats = Reference(ws, min_col=1, min_row=2, max_row=len(df)+1)

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)

    ws.add_chart(chart, "G2")

    wb.save(output_path)

    print("✅ Excel report with chart generated!")

# ✅ Run the script
generate_report("raw_data.csv", "final_report.xlsx")


